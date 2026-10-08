"""Pruebas para Fase 7: Reportes, Bitácoras, Kárdex, Existencias y Exportaciones XLSX."""
import io
import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone

from apps.cases.models import EquipmentCase
from apps.inventory.models import EquipmentModel, InventoryItem, InventoryMovement
from apps.org.models import Center, Region, Zone
from apps.reports.services import (
    export_queryset_to_xlsx,
    get_kardex_report_data,
    get_stock_balances_summary,
    get_vouchers_report_data,
    get_workshop_cases_report_data,
)
from apps.transitions.services import sync_transitions_to_db
from apps.vouchers.models import Voucher, VoucherLine

User = get_user_model()


@pytest.fixture(autouse=True)
def setup_reports_data(db):
    call_command("seed_roles")
    call_command("seed_catalogs")
    call_command("seed_org")
    call_command("seed_demo")
    sync_transitions_to_db()


@pytest.fixture
def inventory_sample(db):
    reg_oaxc = Region.objects.get(code="OAXC")
    reg_other = Region.objects.exclude(code="OAXC").first()
    u_gerente = User.objects.get(email="gerente.ssit@test-ssit.local")
    u_almacen = User.objects.get(email="almacen.oaxc@test-ssit.local")

    model_a = EquipmentModel.objects.first()
    model_b = EquipmentModel.objects.last()

    now = timezone.now()

    # Movimiento 1: Entrada
    m1 = InventoryMovement.objects.create(
        equipment_model=model_a,
        quantity=5,
        movement_type=InventoryMovement.MovementType.ENTRADA,
        region=reg_oaxc,
        warehouse=InventoryItem.Warehouse.STOCK,
        condition=InventoryItem.Condition.NUEVO,
        serial_number="SN-TEST-001",
        moved_at=now,
        confirmed=True,
        created_by=u_almacen,
    )
    # Movimiento 2: Salida
    m2 = InventoryMovement.objects.create(
        equipment_model=model_a,
        quantity=2,
        movement_type=InventoryMovement.MovementType.SALIDA,
        region=reg_oaxc,
        warehouse=InventoryItem.Warehouse.STOCK,
        condition=InventoryItem.Condition.NUEVO,
        serial_number="SN-TEST-001",
        moved_at=now,
        confirmed=True,
        created_by=u_almacen,
    )
    # Movimiento 3: Entrada en otra región
    m3 = InventoryMovement.objects.create(
        equipment_model=model_b,
        quantity=10,
        movement_type=InventoryMovement.MovementType.ENTRADA,
        region=reg_other,
        warehouse=InventoryItem.Warehouse.PROYECTOS,
        condition=InventoryItem.Condition.RECUPERADO,
        serial_number="SN-TEST-002",
        moved_at=now,
        confirmed=True,
        created_by=u_gerente,
    )
    return {
        "reg_oaxc": reg_oaxc,
        "reg_other": reg_other,
        "u_gerente": u_gerente,
        "u_almacen": u_almacen,
        "model_a": model_a,
        "model_b": model_b,
        "m1": m1,
        "m2": m2,
        "m3": m3,
    }


@pytest.mark.django_db
class TestReportsAndKardexServices:
    def test_export_queryset_to_xlsx(self):
        rows = [
            {"id": 1, "name": "Equipo 1", "qty": 10},
            {"id": 2, "name": "Equipo 2", "qty": 5},
        ]
        columns = [("id", "ID"), ("name", "Nombre"), ("qty", "Cantidad")]
        buffer = export_queryset_to_xlsx(rows, columns, sheet_title="TestSheet")
        assert isinstance(buffer, io.BytesIO)

        # Cargar el libro Excel generado con openpyxl para verificar estructura
        wb = openpyxl.load_workbook(buffer)
        assert "TestSheet" in wb.sheetnames
        ws = wb["TestSheet"]
        assert ws.cell(row=1, column=1).value == "ID"
        assert ws.cell(row=1, column=2).value == "Nombre"
        assert ws.cell(row=1, column=3).value == "Cantidad"
        assert ws.cell(row=2, column=2).value == "Equipo 1"
        assert ws.cell(row=3, column=3).value == 5

    def test_get_kardex_report_data_scope_and_filters(self, inventory_sample):
        # Usuario regional solo ve su región
        qs_regional = get_kardex_report_data(user=inventory_sample["u_almacen"])
        assert all(m.region == inventory_sample["reg_oaxc"] for m in qs_regional)

        # Usuario global ve todo
        qs_global = get_kardex_report_data(user=inventory_sample["u_gerente"])
        assert qs_global.count() >= 3

        # Filtro por número de serie
        qs_serial = get_kardex_report_data(user=inventory_sample["u_gerente"], serial_number="SN-TEST-001")
        assert qs_serial.count() == 2

        # Filtro por almacén
        qs_wh = get_kardex_report_data(user=inventory_sample["u_gerente"], warehouse="PROYECTOS")
        assert qs_wh.count() >= 1

        # Filtro por modelo y condición
        qs_model = get_kardex_report_data(
            user=inventory_sample["u_gerente"],
            model_id=inventory_sample["model_a"].id,
            condition=InventoryItem.Condition.NUEVO,
        )
        assert qs_model.count() == 2

        # Filtro por fechas y región
        today_str = timezone.now().strftime("%Y-%m-%d")
        qs_dates = get_kardex_report_data(
            user=inventory_sample["u_gerente"],
            start_date="2020-01-01",
            end_date="2030-12-31",
            region_id=inventory_sample["reg_oaxc"].id,
        )
        assert qs_dates.count() == 2

        # Usuario no autenticado
        assert get_kardex_report_data(None).count() == 0

    def test_get_stock_balances_summary_filters(self, inventory_sample):
        # Saldo para OAXC (5 - 2 = 3 para model_a)
        balances_regional = get_stock_balances_summary(
            inventory_sample["u_almacen"],
            region_id=inventory_sample["reg_oaxc"].id,
            warehouse=InventoryItem.Warehouse.STOCK,
        )
        match = [b for b in balances_regional if b["equipment_model__id"] == inventory_sample["model_a"].id]
        assert len(match) == 1
        assert match[0]["total_balance"] == 3

        # Usuario nulo
        assert get_stock_balances_summary(None) == []

    def test_vouchers_and_workshop_services_filters(self, inventory_sample):
        # Vouchers service filters
        v_qs = get_vouchers_report_data(
            user=inventory_sample["u_gerente"],
            start_date="2020-01-01",
            end_date="2030-12-31",
            region_id=inventory_sample["reg_oaxc"].id,
            status=Voucher.Status.SURTIDO,
            origin_type=Voucher.OriginType.VALE,
        )
        assert get_vouchers_report_data(None).count() == 0

        # Workshop service filters
        w_qs = get_workshop_cases_report_data(
            user=inventory_sample["u_gerente"],
            start_date="2020-01-01",
            end_date="2030-12-31",
            workshop_region_id=inventory_sample["reg_other"].id,
            stage="TALLER_HABILITADO",
            engineer_id=inventory_sample["u_gerente"].id,
        )
        assert get_workshop_cases_report_data(None).count() == 0


@pytest.mark.django_db
class TestReportsViews:
    def test_report_dashboard_access(self, client):
        u_gerente = User.objects.get(email="gerente.ssit@test-ssit.local")
        u_almacen = User.objects.get(email="almacen.oaxc@test-ssit.local")

        client.force_login(u_gerente)
        url = reverse("reports:dashboard")
        resp = client.get(url)
        assert resp.status_code == 200
        assert "Centro de Bitácoras y Reportes" in resp.content.decode("utf-8")

        client.force_login(u_almacen)
        resp = client.get(url)
        assert resp.status_code == 200

    def test_kardex_list_view_and_xlsx_export(self, client, inventory_sample):
        client.force_login(inventory_sample["u_almacen"])
        url = reverse("reports:kardex_list")

        # Vista HTML
        resp = client.get(url)
        assert resp.status_code == 200
        assert "Kárdex de Movimientos" in resp.content.decode("utf-8")
        assert "SN-TEST-001" in resp.content.decode("utf-8")

        # Exportación XLSX
        resp_xlsx = client.get(f"{url}?export=xlsx")
        assert resp_xlsx.status_code == 200
        assert resp_xlsx["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert 'attachment; filename="kardex_movimientos.xlsx"' in resp_xlsx["Content-Disposition"]

    def test_stock_balance_view_and_xlsx_export(self, client, inventory_sample):
        client.force_login(inventory_sample["u_gerente"])
        url = reverse("reports:stock_balance")

        # Vista HTML
        resp = client.get(url)
        assert resp.status_code == 200
        assert "Existencias y Saldos de Almacén" in resp.content.decode("utf-8")

        # Exportación XLSX
        resp_xlsx = client.get(f"{url}?export=xlsx")
        assert resp_xlsx.status_code == 200
        assert resp_xlsx["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_vouchers_report_view_and_export(self, client, inventory_sample):
        reg = inventory_sample["reg_oaxc"]
        center = Center.objects.filter(region=reg).first()
        voucher = Voucher.objects.create(
            origin_type=Voucher.OriginType.VALE,
            report_folio="REP-VALE-101",
            region=reg,
            destination_center=center,
            status=Voucher.Status.SURTIDO,
            requested_by=inventory_sample["u_almacen"],
            delivered_at=timezone.now(),
        )

        client.force_login(inventory_sample["u_gerente"])
        url = reverse("reports:vouchers_report")

        resp = client.get(url)
        assert resp.status_code == 200
        assert "REP-VALE-101" in resp.content.decode("utf-8")

        resp_xlsx = client.get(f"{url}?export=xlsx")
        assert resp_xlsx.status_code == 200
        assert resp_xlsx["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_workshop_report_view_and_export(self, client, inventory_sample):
        case = EquipmentCase.objects.create(
            origin_region=inventory_sample["reg_oaxc"],
            workshop_region=inventory_sample["reg_other"],
            returned_model=inventory_sample["model_a"],
            returned_serial="SN-WS-99",
            stage="TALLER_HABILITADO",
            definition=EquipmentCase.Definition.REPARADO,
            assigned_engineer=inventory_sample["u_gerente"],
            repair_folio="FOL-REP-555",
            received_by=inventory_sample["u_almacen"],
        )

        client.force_login(inventory_sample["u_gerente"])
        url = reverse("reports:workshop_report")

        resp = client.get(url)
        assert resp.status_code == 200
        assert "SN-WS-99" in resp.content.decode("utf-8")
        assert "FOL-REP-555" in resp.content.decode("utf-8")

        resp_xlsx = client.get(f"{url}?export=xlsx")
        assert resp_xlsx.status_code == 200
        assert resp_xlsx["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_zone_summary_report_access(self, client, inventory_sample):
        client.force_login(inventory_sample["u_gerente"])
        url = reverse("reports:zone_summary")

        resp = client.get(url)
        assert resp.status_code == 200
        assert "Resumen Consolidado por Zona" in resp.content.decode("utf-8")

        resp_xlsx = client.get(f"{url}?export=xlsx")
        assert resp_xlsx.status_code == 200
        assert resp_xlsx["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_coupa_report_permissions(self, client, inventory_sample):
        reg = inventory_sample["reg_oaxc"]
        center = Center.objects.filter(region=reg).first()
        v_coupa = Voucher.objects.create(
            origin_type=Voucher.OriginType.COUPA,
            coupa_id="COUPA-7788",
            report_folio="COUPA-REQ-7788",
            region=reg,
            destination_center=center,
            status=Voucher.Status.SURTIDO,
            requested_by=inventory_sample["u_almacen"],
        )

        url = reverse("reports:coupa_report")

        # Soporte regional sin permiso 'view_coupa_consumption'
        u_soporte = User.objects.get(email="soporte.oaxc@test-ssit.local")
        client.force_login(u_soporte)
        resp_unauth = client.get(url)
        assert resp_unauth.status_code == 302  # Redirige al dashboard

        # Gerente SSIT SÍ tiene permiso
        client.force_login(inventory_sample["u_gerente"])
        resp_auth = client.get(url)
        assert resp_auth.status_code == 200
        assert "COUPA-REQ-7788" in resp_auth.content.decode("utf-8")

        # Exportación XLSX
        resp_xlsx = client.get(f"{url}?export=xlsx")
        assert resp_xlsx.status_code == 200
        assert resp_xlsx["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
