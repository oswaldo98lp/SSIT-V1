"""ETL and Inventory Reconciliation Services for SSIT 2.0."""
import csv
import io
from datetime import datetime
from typing import Any

from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import Case, F, IntegerField, Sum, Value, When
from django.utils import timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.cases.models import EquipmentCase
from apps.catalog.models import CatalogValue
from apps.inventory.models import EquipmentModel, InventoryItem, InventoryMovement
from apps.org.models import Center, Region
from apps.vouchers.models import Voucher, VoucherLine

User = get_user_model()


def import_legacy_vouchers_from_data(rows: list[dict[str, Any]], dry_run: bool = False) -> dict[str, Any]:
    """
    Imports legacy vouchers from parsed dictionary records.
    Idempotent using legacy_id.
    """
    stats = {
        "total": len(rows),
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "errors": [],
    }

    system_user = User.objects.filter(is_superuser=True).first() or User.objects.first()

    for idx, row in enumerate(rows, start=1):
        legacy_id = str(row.get("legacy_id") or "").strip()
        if not legacy_id:
            stats["skipped"] += 1
            stats["errors"].append(f"Fila {idx}: 'legacy_id' es obligatorio.")
            continue

        try:
            with transaction.atomic():
                # 1. Resolver Región
                region_code = str(row.get("region_code") or "").strip().upper()
                region = Region.objects.filter(code=region_code).first()
                if not region:
                    region = Region.objects.first()

                # 2. Resolver Centro
                center_val = str(row.get("center_number") or row.get("center_code") or "").strip()
                center = None
                if center_val.isdigit():
                    center = Center.objects.filter(number=int(center_val)).first()
                if not center and center_val:
                    center = Center.objects.filter(name__icontains=center_val).first()
                if not center and region:
                    center = Center.objects.filter(region=region).first()

                # 3. Resolver Solicitante
                user_email = str(row.get("user_email") or "").strip().lower()
                requested_by = User.objects.filter(email=user_email).first() or system_user

                # 4. Tipo de Origen y Coupa ID
                origin_raw = str(row.get("origin_type") or "VALE").strip().upper()
                origin_type = Voucher.OriginType.COUPA if "COUPA" in origin_raw else Voucher.OriginType.VALE
                coupa_id = str(row.get("coupa_id") or "").strip() or None
                if origin_type == Voucher.OriginType.COUPA and not coupa_id:
                    coupa_id = f"COUPA-LEGACY-{legacy_id}"

                # 5. Estatus
                status_raw = str(row.get("status") or "SURTIDO").strip().upper()
                if "CANCEL" in status_raw:
                    status = Voucher.Status.CANCELADO
                elif "SURTIDO" in status_raw or "ENTREGADO" in status_raw:
                    status = Voucher.Status.SURTIDO
                else:
                    status = Voucher.Status.PENDIENTE_SURTIDO

                report_folio = str(row.get("report_folio") or "").strip() or f"FOL-LEGACY-{legacy_id}"
                comments = str(row.get("comments") or "").strip()

                voucher, created = Voucher.objects.update_or_create(
                    legacy_id=legacy_id,
                    defaults={
                        "region": region,
                        "destination_center": center,
                        "requested_by": requested_by,
                        "origin_type": origin_type,
                        "coupa_id": coupa_id,
                        "status": status,
                        "report_folio": report_folio,
                        "comments": comments,
                    },
                )

                if created:
                    stats["created"] += 1
                else:
                    stats["updated"] += 1

                if dry_run:
                    transaction.set_rollback(True)

        except Exception as exc:
            stats["skipped"] += 1
            stats["errors"].append(f"Fila {idx} (ID {legacy_id}): {str(exc)}")

    return stats


def import_legacy_cases_from_data(rows: list[dict[str, Any]], dry_run: bool = False) -> dict[str, Any]:
    """
    Imports legacy equipment return cases.
    Idempotent using legacy_id.
    """
    stats = {
        "total": len(rows),
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "errors": [],
    }

    system_user = User.objects.filter(is_superuser=True).first() or User.objects.first()

    for idx, row in enumerate(rows, start=1):
        legacy_id = str(row.get("legacy_id") or "").strip()
        if not legacy_id:
            stats["skipped"] += 1
            stats["errors"].append(f"Fila {idx}: 'legacy_id' es obligatorio.")
            continue

        try:
            with transaction.atomic():
                origin_code = str(row.get("origin_region_code") or "").strip().upper()
                origin_region = Region.objects.filter(code=origin_code).first() or Region.objects.first()

                workshop_code = str(row.get("workshop_region_code") or "").strip().upper()
                workshop_region = Region.objects.filter(code=workshop_code).first()

                returned_model_name = str(row.get("returned_model_name") or "").strip()
                returned_model = EquipmentModel.objects.filter(name__icontains=returned_model_name).first()

                returned_serial = str(row.get("returned_serial") or "").strip() or f"SN-LEGACY-{legacy_id}"
                repair_folio = str(row.get("repair_folio") or "").strip()
                stage = str(row.get("stage") or "CONFIRMADO").strip().upper()

                def_raw = str(row.get("definition") or "").strip().upper()
                if "REPARADO" in def_raw:
                    definition = EquipmentCase.Definition.REPARADO
                elif "HUESARIO" in def_raw or "SCRAP" in def_raw:
                    definition = EquipmentCase.Definition.HUESARIO
                elif "GARANTIA" in def_raw:
                    definition = EquipmentCase.Definition.GARANTIA
                else:
                    definition = None

                user_email = str(row.get("received_by_email") or "").strip().lower()
                received_by = User.objects.filter(email=user_email).first() or system_user

                case, created = EquipmentCase.objects.update_or_create(
                    legacy_id=legacy_id,
                    defaults={
                        "origin_region": origin_region,
                        "workshop_region": workshop_region,
                        "returned_model": returned_model,
                        "returned_model_text": returned_model_name if not returned_model else "",
                        "returned_serial": returned_serial,
                        "repair_folio": repair_folio,
                        "stage": stage,
                        "definition": definition,
                        "received_by": received_by,
                    },
                )

                if created:
                    stats["created"] += 1
                else:
                    stats["updated"] += 1

                if dry_run:
                    transaction.set_rollback(True)

        except Exception as exc:
            stats["skipped"] += 1
            stats["errors"].append(f"Fila {idx} (ID {legacy_id}): {str(exc)}")

    return stats


def import_legacy_inventory_items_from_data(rows: list[dict[str, Any]], dry_run: bool = False) -> dict[str, Any]:
    """
    Imports physical inventory items and registers initial stock entries in Kardex.
    """
    stats = {
        "total": len(rows),
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "errors": [],
    }

    system_user = User.objects.filter(is_superuser=True).first() or User.objects.first()
    now = timezone.now()

    for idx, row in enumerate(rows, start=1):
        serial_number = str(row.get("serial_number") or "").strip()
        model_name = str(row.get("model_name") or "").strip()
        region_code = str(row.get("region_code") or "").strip().upper()

        if not serial_number or not model_name:
            stats["skipped"] += 1
            stats["errors"].append(f"Fila {idx}: 'serial_number' y 'model_name' son obligatorios.")
            continue

        try:
            with transaction.atomic():
                region = Region.objects.filter(code=region_code).first() or Region.objects.first()
                equipment_model, _ = EquipmentModel.objects.get_or_create(
                    name=model_name,
                    defaults={"brand": str(row.get("brand") or "Genérico").strip(), "is_inventoriable": True},
                )

                condition_raw = str(row.get("condition") or "NUEVO").strip().upper()
                condition = InventoryItem.Condition.RECUPERADO if "RECUP" in condition_raw else InventoryItem.Condition.NUEVO

                warehouse_raw = str(row.get("warehouse") or "STOCK").strip().upper()
                warehouse = InventoryItem.Warehouse.PROYECTOS if "PROY" in warehouse_raw else InventoryItem.Warehouse.STOCK

                voucher_legacy_id = str(row.get("voucher_legacy_id") or "").strip()
                voucher = Voucher.objects.filter(legacy_id=voucher_legacy_id).first() if voucher_legacy_id else None

                if not voucher:
                    # Crear vale contenedor de saldo inicial si no existe
                    center = Center.objects.filter(region=region).first()
                    voucher, _ = Voucher.objects.get_or_create(
                        report_folio="INICIAL-MIGRACION",
                        region=region,
                        defaults={
                            "destination_center": center,
                            "requested_by": system_user,
                            "origin_type": Voucher.OriginType.VALE,
                            "status": Voucher.Status.SURTIDO,
                            "comments": "Carga inicial de inventario histórico",
                        },
                    )

                line, _ = VoucherLine.objects.get_or_create(
                    voucher=voucher,
                    equipment_model=equipment_model,
                    defaults={"quantity": 1, "status": VoucherLine.LineStatus.SURTIDO},
                )

                item, created = InventoryItem.objects.update_or_create(
                    serial_number=serial_number,
                    defaults={
                        "voucher": voucher,
                        "line": line,
                        "equipment_model": equipment_model,
                        "region": region,
                        "condition": condition,
                        "warehouse": warehouse,
                    },
                )

                # Registrar movimiento de entrada en Kárdex si es nuevo
                if created:
                    InventoryMovement.objects.create(
                        equipment_model=equipment_model,
                        movement_type=InventoryMovement.MovementType.ENTRADA,
                        serial_number=serial_number,
                        condition=condition,
                        warehouse=warehouse,
                        region=region,
                        quantity=1,
                        moved_at=now,
                        transfer_folio=f"INICIAL-{serial_number}",
                        item=item,
                        confirmed=True,
                        created_by=system_user,
                    )
                    stats["created"] += 1
                else:
                    stats["updated"] += 1

                if dry_run:
                    transaction.set_rollback(True)

        except Exception as exc:
            stats["skipped"] += 1
            stats["errors"].append(f"Fila {idx} ({serial_number}): {str(exc)}")

    return stats


def reconcile_inventory_balances(physical_records: list[dict[str, Any]]) -> dict[str, Any]:
    """
    Reconciles physical inventory records against computed Kardex movements.
    Detects surpluses, shortages, and matched balances.
    """
    reconciliation_lines = []
    total_physical = 0
    total_system = 0
    matched_count = 0
    surplus_count = 0
    shortage_count = 0

    # Cache de saldos del sistema
    movements_qs = InventoryMovement.objects.filter(confirmed=True).values(
        "equipment_model_id",
        "equipment_model__name",
        "equipment_model__brand",
        "region_id",
        "region__code",
        "region__name",
        "warehouse",
        "condition",
    ).annotate(
        system_balance=Sum(
            Case(
                When(movement_type=InventoryMovement.MovementType.ENTRADA, then="quantity"),
                When(movement_type=InventoryMovement.MovementType.SALIDA, then=-1 * F("quantity")),
                default=Value(0),
                output_field=IntegerField(),
            )
        )
    )

    system_lookup = {}
    for m in movements_qs:
        key = (
            m["equipment_model_id"],
            m["region_id"],
            m["warehouse"],
            m["condition"],
        )
        system_lookup[key] = m

    processed_keys = set()

    for idx, row in enumerate(physical_records, start=1):
        model_name = str(row.get("model_name") or "").strip()
        region_code = str(row.get("region_code") or "").strip().upper()
        warehouse = str(row.get("warehouse") or "STOCK").strip().upper()
        condition = str(row.get("condition") or "NUEVO").strip().upper()
        try:
            physical_qty = int(row.get("physical_quantity", 0))
        except (ValueError, TypeError):
            physical_qty = 0

        # Resolver modelo y región
        model = EquipmentModel.objects.filter(name__icontains=model_name).first()
        region = Region.objects.filter(code=region_code).first()

        model_id = model.id if model else None
        region_id = region.id if region else None
        key = (model_id, region_id, warehouse, condition)
        processed_keys.add(key)

        system_entry = system_lookup.get(key)
        system_qty = system_entry["system_balance"] if system_entry else 0

        difference = physical_qty - system_qty
        if difference == 0:
            status = "COINCIDE"
            matched_count += 1
        elif difference > 0:
            status = "SOBRANTE"
            surplus_count += 1
        else:
            status = "FALTANTE"
            shortage_count += 1

        total_physical += physical_qty
        total_system += system_qty

        reconciliation_lines.append({
            "line_number": idx,
            "model_name": model.name if model else model_name,
            "brand": model.brand if model else "-",
            "region_code": region.code if region else region_code,
            "warehouse": warehouse,
            "condition": condition,
            "physical_quantity": physical_qty,
            "system_quantity": system_qty,
            "difference": difference,
            "status": status,
        })

    total_lines = len(reconciliation_lines)
    match_rate = round((matched_count / total_lines * 100), 2) if total_lines > 0 else 100.0

    return {
        "summary": {
            "total_lines": total_lines,
            "matched_count": matched_count,
            "surplus_count": surplus_count,
            "shortage_count": shortage_count,
            "total_physical": total_physical,
            "total_system": total_system,
            "match_rate": match_rate,
            "reconciled_at": timezone.now(),
        },
        "lines": reconciliation_lines,
    }


def generate_reconciliation_xlsx(reconciliation_data: dict[str, Any]) -> io.BytesIO:
    """
    Exports reconciliation results to a styled Excel (XLSX) workbook.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conciliacion_Inventario"

    # Estilos
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # Verde claro
    yellow_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")  # Amarillo claro
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # Rojo claro

    # Encabezados
    headers = [
        ("line_number", "#"),
        ("model_name", "Artículo / Modelo"),
        ("brand", "Marca"),
        ("region_code", "Región"),
        ("warehouse", "Almacén"),
        ("condition", "Condición"),
        ("physical_quantity", "Físico Reportado"),
        ("system_quantity", "Saldo Kárdex"),
        ("difference", "Diferencia"),
        ("status", "Estatus Conciliación"),
    ]

    ws.row_dimensions[1].height = 26
    for col_idx, (_, header_label) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Filas de datos
    for row_idx, item in enumerate(reconciliation_data.get("lines", []), start=2):
        ws.row_dimensions[row_idx].height = 20
        status = item.get("status")

        for col_idx, (field_key, _) in enumerate(headers, start=1):
            val = item.get(field_key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

            if field_key in ["physical_quantity", "system_quantity", "difference"]:
                cell.alignment = right_align
                cell.font = bold_font
            elif field_key in ["line_number", "region_code", "warehouse", "condition", "status"]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Color según estatus
            if status == "COINCIDE":
                cell.fill = green_fill
            elif status == "SOBRANTE":
                cell.fill = yellow_fill
            elif status == "FALTANTE":
                cell.fill = red_fill

    # Ajuste automático de ancho
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
